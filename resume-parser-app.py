
import re
from pathlib import Path

import joblib
import PyPDF2
import streamlit as st
from docx import Document

# 1. PAGE CONFIGURATION
st.set_page_config(
    page_title="Resume Parsing and Job Recommendation System",
    page_icon="📄",
    layout="wide"
)

st.title("Resume Parsing, Job Categorization and Recommendation System")

st.write(
    "Upload a resume to extract candidate information, "
    "education, skills, work experience, job category "
    "and job recommendations."
)


# 2. LOAD TRAINED MODELS
@st.cache_resource
def load_models():
    vectorizer_cat= joblib.load( "vectorizer_job_cat.joblib")
    vectorizer_recom = joblib.load("vectorizer_job_recom.joblib")
    category_model = joblib.load("model_job_cat.joblib")
    # Trained job recommendation model
    recommendation_model = joblib.load("model_job_recom.joblib")
    return (
        vectorizer_cat,
        vectorizer_recom,
        category_model,
        recommendation_model
    )
try:
    (
        vectorizer_cat,
        vectorizer_recom,
        category_model,
        recommendation_model
    ) = load_models()

except Exception as e:
    st.error(f"Unable to load the trained models: {e}")
    st.stop()
# 3. EXTRACT TEXT FROM PDF
def extract_text_from_pdf(uploaded_file):
    text = ""
    pdf_reader = PyPDF2.PdfReader(uploaded_file)
    for page in pdf_reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text
# 4. EXTRACT TEXT FROM DOCX
def extract_text_from_docx(uploaded_file):
    document = Document(uploaded_file)
    text = []
    # Extract paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())
    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if row_text:
                text.append(" | ".join(row_text))

    return "\n".join(text)


# 5. EXTRACT RESUME TEXT

def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return ""
# 6. CLEAN TEXT FOR THE ML MODELS
def cleanResume(txt):
    txt = re.sub(r"http\S+|www\S+", " ", txt)
    txt = re.sub(r"\bRT\b|\bcc\b", " ", txt)
    txt = re.sub(r"@\S+", " ", txt)
    txt = re.sub(r"#(\S+)", r"\1", txt)
    txt = re.sub(r"[^\x00-\x7F]", " ", txt)
    txt = re.sub(
        r"""[!"#$%&'()*+,\-./:;<=>?@\[\]^_`{|}~]""",
        " ",
        txt
    )
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt
# 7. EXTRACT PHONE NUMBER
def extract_phone_number(text):
    pattern = (
        r"(?<!\d)(?:\+234|234|0)[\s-]?"
        r"(?:7|8|9)\d{2}[\s-]?\d{3}[\s-]?\d{4}(?!\d)"
    )
    matches = re.findall(pattern, text)
    if not matches:
        return None
    phone = re.sub(r"[\s-]", "", matches[0])
    if phone.startswith("0"):
        phone = "+234" + phone[1:]

    elif phone.startswith("234"):
        phone = "+" + phone
    return phone
# 8. EXTRACT EMAIL
def extract_email(text):
    pattern = (
        r"[A-Za-z0-9._%+-]+@"
        r"[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
    )
    match = re.search(pattern, text)
    return match.group(0) if match else None
# 9. SKILLS LIST

skills = [
    # Programming Languages
    "Python",
    "Java",
    "JavaScript",
    "C++",
    "C#",
    "Ruby",
    "Swift",
    "Kotlin",
    "SQL",
    "Bash",

    # Web Development
    "HTML",
    "CSS",
    "React",
    "Angular",
    "Vue.js",
    "Node.js",
    "Django",
    "Spring",
    "Express",
    "JSP",
    "Servlets",
    "React Native",
    "Flutter",

    # Data Science / Machine Learning
    "Machine Learning",
    "Deep Learning",
    "Data Science",
    "Data Analysis",
    "Data Visualization",
    "Data Modeling",
    "Data Cleansing",
    "Statistical Analysis",
    "Predictive Modeling",
    "Machine Learning Algorithms",
    "TensorFlow",
    "PyTorch",
    "Scikit-learn",
    "Pandas",
    "NumPy",
    "Big Data",

    # Databases
    "MySQL",
    "Oracle",
    "SQL Server",
    "PostgreSQL",
    "MongoDB",
    "Cassandra",
    "Database Management",
    "Database Administration",
    "Database Security",
    "Database Tuning",
    "Data Warehousing",

    # Big Data
    "Hadoop",
    "Spark",
    "Informatica",
    "Talend",

    # Cloud / DevOps
    "AWS",
    "Azure",
    "GCP",
    "Docker",
    "Kubernetes",
    "Terraform",
    "Ansible",
    "CI/CD",
    "Git",
    "GitHub",
    "Jira",
    "Cloud Computing",
    "Server Management",
    "Server Maintenance",

    # Cybersecurity / Networking
    "Network Security",
    "Network Management",
    "IT Security",
    "Cybersecurity",
    "Intrusion Detection",
    "Incident Response",
    "SIEM",
    "IDS/IPS",
    "TCP/IP",
    "CCNA",
    "Wireshark",
    "VMware",
    "Hyper-V",
    "Active Directory",
    "Troubleshooting",

    # Testing / QA
    "Quality Assurance",
    "Software Quality Assurance",
    "Software Testing",
    "Manual Testing",
    "Automated Testing",
    "Test Automation",
    "Test Planning",
    "Test Case Design",
    "Test Execution",
    "Bug Tracking",
    "Selenium",
    "Appium",
    "JMeter",

    # Business / Management
    "Project Management",
    "Product Management",
    "Business Management",
    "Business Development",
    "Strategic Planning",
    "Team Leadership",
    "Team Management",
    "Team Coordination",
    "Risk Management",
    "Risk Assessment",
    "Change Management",
    "Performance Management",
    "Process Management",
    "Operations Management",

    # Finance / Accounting
    "Financial Analysis",
    "Financial Planning",
    "Financial Management",
    "Investment Analysis",
    "Investment Strategy",
    "Portfolio Management",
    "Retirement Planning",
    "Tax Planning",
    "Tax Compliance",
    "Tax Law",
    "Tax Preparation",
    "Financial Reporting",
    "Budget Management",
    "Budgeting",
    "Cost Analysis",
    "QuickBooks",

    # Marketing
    "Digital Marketing",
    "Content Marketing",
    "Content Creation",
    "Content Strategy",
    "Social Media Marketing",
    "Social Media Management",
    "Social Media Analytics",
    "Email Marketing",
    "Marketing Analytics",
    "Marketing Automation",
    "SEO",
    "SEM",
    "Search Engine Optimization",
    "Keyword Research",
    "Google Analytics",
    "Facebook Insights",
    "Semrush",
    "Ahrefs",
    "Screaming Frog",
    "Mailchimp",
    "HubSpot",
    "Marketo",
    "Salesforce",

    # Sales / Customer
    "Sales Strategy",
    "Sales Management",
    "Sales Techniques",
    "Business Development",
    "Account Management",
    "Customer Relationship Management",
    "Client Relationship Management",
    "Customer Service",
    "Customer Support",
    "Customer Experience",
    "Customer Advocacy",
    "Relationship Building",
    "Negotiation",
    "Contract Negotiation",
    "Cross-selling",
    "Upselling",

    # Supply Chain
    "Supply Chain Management",
    "Supply Chain Logistics",
    "Procurement",
    "Procurement Management",
    "Procurement Processes",
    "Supplier Management",
    "Vendor Management",
    "Vendor Coordination",
    "Vendor Assessment",
    "Inventory Management",
    "Inventory Control",
    "Demand Forecasting",
    "Purchase Order Management",
    "Spend Analysis",

    # Design
    "Graphic Design",
    "UI Design",
    "UX Design",
    "User Interface Design",
    "User Experience Design",
    "User-Centered Design",
    "Responsive Design",
    "Web Design",
    "CAD",
    "CAD/CAM",
    "AutoCAD",
    "AutoCAD Civil 3D",
    "SolidWorks",
    "Revit",
    "MicroStation",
    "Sketch",
    "Figma",
    "Adobe Photoshop",
    "Adobe Illustrator",
    "Adobe XD",
    "Adobe InDesign",
    "Adobe Creative Suite",
    "SketchUp",

    # Data / Business Intelligence
    "Business Intelligence",
    "Data Interpretation",
    "Data Reporting",
    "Tableau",
    "Power BI",
    "Excel",
    "SPSS",
    "SAS",
    "Stata",
    "NVivo",
    "Microsoft Access",

    # Research
    "Market Research",
    "Research Methods",
    "Data Collection",
    "Survey Design",
    "Survey Analysis",
    "Consumer Behavior Analysis",
    "Competitor Analysis",
    "Trend Analysis",
    "Qualitative Research",
    "Quantitative Research",
    "Ethnographic Research",
    "Statistical Research",

    # Administration
    "Office Administration",
    "Office Management",
    "Calendar Management",
    "Meeting Scheduling",
    "Travel Coordination",
    "Document Management",
    "Record Keeping",
    "Record Management",
    "Time Management",
    "Event Planning",
    "Event Management",

    # Communication
    "Communication",
    "Written Communication",
    "Verbal Communication",
    "Public Speaking",
    "Presentation Skills",
    "Problem Solving",
    "Critical Thinking",
    "Decision Making",
    "Attention to Detail",
    "Leadership",
    "Teamwork",
    "Collaboration",
    "Creativity",

    # Legal
    "Regulatory Compliance",
    "Legal Research",
    "Contract Law",
    "Tax Law",
    "Corporate Law",
    "Litigation",
    "E-Discovery",
    "LexisNexis",
    "Relativity",

    # Healthcare
    "Patient Care",
    "Patient Education",
    "Clinical Assessment",
    "Health Assessment",
    "Critical Care",
    "Geriatric Care",
    "Pediatric Care",
    "Mental Health",
    "Addiction Recovery",
    "Motivational Interviewing",
    "Relapse Prevention",

    # Engineering
    "Engineering Design",
    "Computer-Aided Design",
    "CAD",
    "CAD/CAM",
    "Avionics",
    "Controls Engineering",
    "Construction Management",
    "Structural Engineering",
    "Environmental Engineering",
    "Chemical Engineering",
    "SAP",
    "SAP2000",

    # E-commerce
    "E-commerce",
    "Shopify"
]
def extract_skills(resume_text):
    found_skills = []
    for skill in skills:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, resume_text, re.IGNORECASE):
            found_skills.append(skill)

    return found_skills
# 10. FIELDS OF STUDY
fields_of_study = [
    "computer science", "computer engineering",
    "software engineering", "information technology",
    "information systems", "data science", "data analytics",
    "artificial intelligence", "cybersecurity",
    "business administration", "accounting", "finance",
    "economics", "marketing", "human resources",
    "mechanical engineering", "electrical engineering",
    "civil engineering", "chemical engineering",
    "industrial engineering", "medicine", "nursing", "pharmacy",
    "biology", "biochemistry", "microbiology", "mathematics",
    "statistics", "physics", "chemistry", "psychology",
    "sociology", "political science", "law", "mass communication",
    "journalism", "architecture", "graphic design",
    "environmental science", "public health"
]
# 11. JOB TITLES
work_experience = [
    "software engineer", "software developer", "web developer",
    "frontend developer", "backend developer", "full stack developer",
    "data scientist", "data analyst", "data engineer",
    "machine learning engineer", "devops engineer", "cloud engineer",
    "network engineer", "cybersecurity analyst",
    "database administrator", "project manager", "product manager",
    "business analyst", "business development manager",
    "financial analyst", "accountant", "auditor",
    "marketing manager", "marketing specialist", "sales manager",
    "sales representative", "customer service representative",
    "human resources manager", "hr specialist", "recruiter",
    "mechanical engineer", "civil engineer", "electrical engineer",
    "chemical engineer", "doctor", "physician", "nurse",
    "pharmacist", "teacher", "lecturer", "professor",
    "lawyer", "attorney", "legal assistant", "graphic designer",
    "ui designer", "ux designer", "architect"
]
# 12. EXTRACT CANDIDATE NAME
def extract_candidate_name(resume_text):
    lines = [
        re.sub(r"\s+", " ", line).strip()
        for line in resume_text.splitlines()
    ]
    name_label = re.compile(
        r"^(?:full name|candidate name|name)\s*[:\-]\s*(.+)$",
        re.IGNORECASE
    )
    excluded = {
        "resume", "curriculum vitae", "cv", "profile",
        "contact", "email", "phone", "address", "education",
        "experience", "skills", "summary", "objective",
        "professional summary", "work experience"
    }
    # Check the first lines because the name is usually near the top.
    for line in lines[:12]:
        if not line:
            continue
        labelled = name_label.match(line)
        if labelled:
            candidate = labelled.group(1).strip(" ,|:-")
            if (
                candidate
                and not re.search(r"@|https?://|www\.|\d", candidate)
                and len(candidate.split()) <= 5
            ):
                return candidate
        if re.search(r"@|https?://|www\.|\d", line):
            continue
        candidate = line.strip(" ,|:-")
        words = candidate.split()
        if candidate.lower() in excluded:
            continue
        if len(words) not in (2, 3, 4):
            continue
        if not all(
            re.fullmatch(r"[A-Za-zÀ-ÖØ-öø-ÿ'’-]+", word)
            for word in words
        ):
            continue
        if any(
            term in candidate.lower()
            for term in [
                "engineer", "developer", "manager", "analyst",
                "student", "university", "college", "polytechnic",
                "bachelor", "computer science"
            ]
        ):
            continue
        return candidate
    return "Name not detected"
# 13. EXTRACT EDUCATION QUALIFICATION
QUALIFICATION_RE = re.compile(
    r"\b(?:"
    r"B\s*\.?\s*Sc\.?|B\s*\.?\s*Eng\.?|B\s*\.?\s*Tech\.?|"
    r"B\s*\.?\s*A\.?|B\s*\.?\s*Ed\.?|B\s*\.?\s*N\.?\s*Sc\.?|"
    r"LLB|LL\.B\.?|M\s*\.?\s*Sc\.?|M\s*\.?\s*Eng\.?|"
    r"M\s*\.?\s*Tech\.?|M\s*\.?\s*A\.?|MBA|HND|OND|NCE|PGD|"
    r"Ph\.?\s*D\.?|Bachelor(?:'s)?\s+(?:of\s+)?"
    r"(?:Science|Engineering|Arts|Technology|Education|Law)|"
    r"Master(?:'s)?\s+(?:of\s+)?"
    r"(?:Science|Engineering|Arts|Technology|Business Administration)|"
    r"Higher National Diploma|National Diploma|Postgraduate Diploma"
    r")\b",
    re.IGNORECASE
)
EDUCATION_HEADINGS = {
    "education", "educational background", "academic background",
    "academic qualifications", "education and qualifications",
    "qualifications"
}
SECTION_HEADINGS = {
    "work experience", "professional experience", "employment history",
    "work history", "career history", "experience", "skills",
    "technical skills", "projects", "certifications", "references",
    "awards", "interests", "languages", "professional summary",
    "career objective", "personal profile"
}
def extract_education_qualification(resume_text):
    lines = [
        re.sub(r"\s+", " ", line).strip(" \t|•-")
        for line in resume_text.splitlines()
        if line.strip()
    ]
    education_lines = []
    in_education = False
    for line in lines:
        heading = line.strip(" :").lower()
        if heading in EDUCATION_HEADINGS:
            in_education = True
            continue
        if in_education and heading in SECTION_HEADINGS:
            break
        if in_education:
            education_lines.append(line)
    # If there is no education heading, search the whole resume.
    search_lines = education_lines if education_lines else lines
    for index, line in enumerate(search_lines):
        match = QUALIFICATION_RE.search(line)
        if not match:
            continue
        result = line.strip(" •\t-–|")
        # Example: "BSc" on one line and "Computer Science" on the next.
        if len(result.split()) <= 2 and index + 1 < len(search_lines):
            next_line = search_lines[index + 1].strip(" •\t-–|")
            if (
                len(next_line.split()) <= 8
                and not re.search(
                    r"university|polytechnic|college|institute",
                    next_line,
                    re.IGNORECASE
                )
            ):
                result = f"{result} in {next_line}"
        return re.sub(r"\s+", " ", result).strip(" ,.;:-")
    # Fallback for lines containing a field of study.
    for line in search_lines:
        for field in sorted(fields_of_study, key=len, reverse=True):
            if re.search(
                r"\b" + re.escape(field) + r"\b",
                line,
                re.IGNORECASE
            ):
                return line.strip(" •\t-–|")

    return "Education qualification not detected"
# 14. EXTRACT EDUCATION FIELDS AND JOB TITLES
def extract_education_and_experience(resume_text):
    keywords = []
    for field in sorted(fields_of_study, key=len, reverse=True):
        pattern = r"(?<!\w)" + re.escape(field) + r"(?!\w)"
        if re.search(pattern, resume_text, re.IGNORECASE):
            keywords.append(field.title())
    for job in sorted(work_experience, key=len, reverse=True):
        pattern = r"(?<!\w)" + re.escape(job) + r"(?!\w)"
        if re.search(pattern, resume_text, re.IGNORECASE):
            keywords.append(job.title())
    return list(dict.fromkeys(keywords))
# 15. EXTRACT WORK EXPERIENCE
MONTHS = (
    r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"
    r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|"
    r"Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
)
DATE_TOKEN = rf"(?:{MONTHS}\.?\s+)?(?:19|20)\d{{2}}"
DATE_RANGE_RE = re.compile(
    rf"(?<!\w){DATE_TOKEN}\s*(?:-|–|—|to|until)\s*"
    rf"(?:present|current|ongoing|{DATE_TOKEN})\b",
    re.IGNORECASE
)
DUTY_WORDS = (
    "responsible for", "duties", "achievements",
    "managed", "developed", "worked on", "led a team",
    "assisted", "created", "implemented"
)
def extract_work_experience_details(resume_text):
    lines = [
        re.sub(r"\s+", " ", line).strip(" \t|•-")
        for line in resume_text.splitlines()
        if line.strip()
    ]
    experience_lines = []
    in_experience = False
    for line in lines:
        heading = line.strip(" :").lower()
        if heading in {
            "work experience", "professional experience",
            "employment history", "work history",
            "career history", "experience"
        }:
            in_experience = True
            continue
        if in_experience and (
            heading in EDUCATION_HEADINGS
            or heading in {
                "skills", "technical skills", "projects",
                "certifications", "references", "awards",
                "interests", "languages"
            }
        ):
            break

        if in_experience:
            experience_lines.append(line)

    search_lines = experience_lines if experience_lines else lines
    entries = []
    current = None
    titles = sorted(work_experience, key=len, reverse=True)
    def save_current():
        nonlocal current
        if current:
            entries.append(current)
        current = None
    for line in search_lines:
        date_match = DATE_RANGE_RE.search(line)
        title_match = None
        for title in titles:
            found = re.search(
                r"(?<!\w)" + re.escape(title) + r"(?!\w)",
                line,
                re.IGNORECASE
            )
            if found:
                title_match = (title, found)
                break
        if title_match:
            save_current()
            title, match = title_match
            company_text = line
            if date_match:
                company_text = (
                    line[:date_match.start()]
                    + " "
                    + line[date_match.end():]
                )
            company_text = re.sub(
                r"(?<!\w)" + re.escape(title) + r"(?!\w)",
                " ",
                company_text,
                count=1,
                flags=re.IGNORECASE
            )
            company_text = re.sub(
                r"^\s*(?:at|@)\s+",
                "",
                company_text,
                flags=re.IGNORECASE
            )

            company_text = re.sub(
                r"\s+", " ", company_text
            ).strip(" ,|:-–—")

            current = {
                "job_title": title.title(),
                "company": company_text or "Not detected",
                "employment_dates": (
                    date_match.group(0)
                    if date_match
                    else "Not detected"
                )
            }

            continue

        if current:
            if date_match and current["employment_dates"] == "Not detected":
                current["employment_dates"] = date_match.group(0)
                continue

            if (
                current["company"] == "Not detected"
                and len(line.split()) <= 12
                and not any(word in line.lower() for word in DUTY_WORDS)
            ):
                current["company"] = line.strip(" ,|:-–—")

    save_current()

    # Remove exact duplicate entries.
    unique_entries = []
    seen = set()

    for entry in entries:
        key = (
            entry["job_title"],
            entry["company"],
            entry["employment_dates"]
        )

        if key not in seen:
            seen.add(key)
            unique_entries.append(entry)

    return unique_entries

# 16. STREAMLIT FILE UPLOAD
uploaded_file = st.file_uploader(
    "Upload Resume",
    type=["pdf", "docx"]
)

if uploaded_file is not None:
    st.success(f"Resume uploaded successfully: {uploaded_file.name}")
    if st.button("Analyze Resume", type="primary"):
        with st.spinner("Extracting and analyzing resume..."):
            try:
                resume_text = extract_resume_text(uploaded_file)
            except Exception as e:
                st.error(f"Could not read the resume: {e}")
                st.stop()
            if not resume_text.strip():
                st.error(
                    "No readable text was found. "
                    "Scanned PDFs may require OCR."
                )
                st.stop()
            # Extract candidate information.
            candidate_name = extract_candidate_name(resume_text)
            phone = extract_phone_number(resume_text)
            email = extract_email(resume_text)
            extracted_skills = extract_skills(resume_text)
            education = extract_education_qualification(resume_text)
            experience_details = extract_work_experience_details(
                resume_text
            )
            education_experience = extract_education_and_experience(
                resume_text
            )
            # Prepare text for your existing trained ML models.
            cleaned_text = cleanResume(resume_text)
            try:
                category_vector = vectorizer_cat.transform(
                    [cleaned_text]
                )
                recom_vector = vectorizer_recom.transform(
                    [cleaned_text]
                )
                predicted_category = category_model.predict(
                    category_vector
                )[0]
                category_confidence = None
                if hasattr(category_model, "predict_proba"):
                    probabilities = category_model.predict_proba(
                        category_vector
                    )[0]
                    category_confidence = (
                        float(max(probabilities)) * 100
                    )
                recommendations = recommendation_model.predict(
                    recom_vector
                )
            except Exception as e:
                st.error(
                    "The resume was parsed, but the ML prediction failed. "
                    f"Check that the model and vectorizer files match. Error: {e}"
                )
                st.stop()
        # 17. DISPLAY CANDIDATE INFORMATION
        st.divider()
        st.header("👤 Candidate Information")
        col1, col2, col3 = st.columns(3)
        with col1:
            st.subheader("Full Name")
            st.write(candidate_name)
        with col2:
            st.subheader("Email Address")
            st.write(email or "Not detected")
        with col3:
            st.subheader("Phone Number")
            st.write(phone or "Not detected")
        # 18. DISPLAY EDUCATION
        st.divider()
        st.header("🎓 Education")
        st.write(education)

        # 19. DISPLAY WORK EXPERIENCE
        st.divider()
        st.header("💼 Work Experience")
        if experience_details:
            for index, item in enumerate(
                experience_details, start=1
            ):
                st.markdown(f"**Experience {index}**")
                st.write(
                    f"**Job Title:** {item['job_title']}"
                )
                st.write(
                    f"**Company:** {item['company']}"
                )
                st.write(
                    f"**Employment Dates:** {item['employment_dates']}"
                )
        else:
            st.info(
                "No structured work experience was detected. "
                "Check the extracted resume text below."
            )

        # 20. DISPLAY SKILL
        st.divider()
        st.header("🛠️ Extracted Skills")

        if extracted_skills:
            for skill in extracted_skills:
                st.write(f"• {skill}")
        else:
            st.info("No matching skills were detected.")
        # 21. DISPLAY EDUCATION FIELDS AND JOB TITLES
        st.divider()
        st.header("📌 Detected Fields and Job Titles")
        if education_experience:
            for item in education_experience:
                st.write(f"• {item}")
        else:
            st.info("No matching fields or job titles were detected.")
        # 22. DISPLAY JOB CATEGORY
        st.divider()
        st.header("🎯 Predicted Job Category")
        st.success(str(predicted_category))
        if category_confidence is not None:
            st.metric(
                "Category Prediction Confidence",
                f"{category_confidence:.2f}%"
            )
        else:
            st.caption(
                "The trained model does not provide predict_proba confidence scores."
            )
        # 23. DISPLAY JOB RECOMMENDATIONS
        st.divider()
        st.header("💼 Recommended Jobs")
        recommendation_list = list(recommendations)
        if recommendation_list:
            for recommendation in recommendation_list:
                st.write(f"• {recommendation}")
        else:
            st.info("No job recommendations were generated.")
        # 24. DISPLAY RAW RESUME TEXT
        st.divider()
        with st.expander("📄 View Extracted Resume Text"):
            st.text_area(
                "Resume Text",
                resume_text,
                height=400
            )
